"""Slide helpers shared by every hepattn meeting deck (16:9, dark slate, Trebuchet).

Copied into each deck folder by new_deck.py so a deck can always be rebuilt on its own,
without the skill installed. Use from build_deck.py:

    from deckkit import *
    prs = new_presentation()
    slide_title(prs, "Title", "Subtitle", "footer")
    slide_figure(prs, "Slide title", "01_figure_name", bullets=[("text", ACCENT, True), ...])
    save_deck(prs, OUT)

Every helper writes text runs tagged English (UK). Without an explicit language PowerPoint can
fall back to line-breaking rules that split words mid-way ("precis / ion").
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
FIGS = HERE / "figures"

W, H = Inches(13.333), Inches(7.5)  # 16:9
BG = RGBColor(0x0F, 0x16, 0x20)
INK = RGBColor(0xE8, 0xED, 0xF4)
MUTED = RGBColor(0x8F, 0xA3, 0xBC)
ACCENT = RGBColor(0xFF, 0xB4, 0x54)  # our new work (amber)
GOOD = RGBColor(0x6B, 0xCB, 0x77)    # confirmed / passing (green)
WARN = RGBColor(0xFF, 0x6B, 0x6B)    # failures / broken (red)
BASE = RGBColor(0x5A, 0xC8, 0xFA)    # the reference / baseline (blue)
VIOLET = RGBColor(0xC7, 0x92, 0xEA)
FONT = "Trebuchet MS"

__all__ = ["ACCENT", "BASE", "GOOD", "INK", "MUTED", "VIOLET", "WARN", "new_presentation", "save_deck",
           "slide_bullets", "slide_figure", "slide_section", "slide_title", "slide_two_col"]


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG


def _text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT, spacing=1.0):
    """runs: list of (text, size, colour, bold)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, colour, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = colour
        r.font.bold = bold
        r.font.name = FONT
        r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_UK
    return box


def _rule(slide, top, colour=ACCENT, left=Inches(0.7), width=Inches(2.0)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(4))
    s.fill.solid()
    s.fill.fore_color.rgb = colour
    s.line.fill.background()
    s.shadow.inherit = False


def _number(slide, n):
    box = slide.shapes.add_textbox(Inches(12.2), Inches(6.85), Inches(0.9), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    r.font.name = FONT
    r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_UK


def _blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    return s


def _heading(s, title):
    _text(s, Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.8), [(title, 26, INK, True)])
    _rule(s, Inches(1.12))


def slide_title(prs, title, subtitle, footer):
    s = _blank(prs)
    _text(s, Inches(1.0), Inches(2.3), Inches(11.3), Inches(1.6), [(title, 44, INK, True)])
    _rule(s, Inches(3.95), left=Inches(1.05), width=Inches(2.6))
    _text(s, Inches(1.0), Inches(4.25), Inches(11.3), Inches(1.4), [(subtitle, 21, MUTED, False)])
    _text(s, Inches(1.0), Inches(6.35), Inches(11.3), Inches(0.6), [(footer, 14, ACCENT, False)])
    return s


def slide_section(prs, kicker, title):
    s = _blank(prs)
    _text(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(0.5), [(kicker.upper(), 15, ACCENT, True)])
    _text(s, Inches(1.0), Inches(3.35), Inches(11.3), Inches(1.4), [(title, 36, INK, True)])
    _rule(s, Inches(4.9), left=Inches(1.05), width=Inches(2.0))
    return s


def slide_figure(prs, title, figure, caption=None, bullets=None):
    """A figure from figures/<figure>.png, scaled to fit, with up to ~3 bullets or one caption.

    bullets: list of (text, colour, bold). Leave room: more than three lines falls off the slide.
    """
    s = _blank(prs)
    _heading(s, title)
    img = FIGS / f"{figure}.png"
    if img.exists():
        from PIL import Image

        iw, ih = Image.open(img).size
        avail_w = Inches(11.9)
        avail_h = Inches(4.15) if bullets else (Inches(4.9) if caption else Inches(5.6))
        scale = min(avail_w / iw, avail_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(str(img), int((W - w) / 2), Inches(1.45), w, h)
        bottom = Inches(1.45) + Emu(h)
    else:
        print(f"  WARNING: missing figure {img.name}")
        bottom = Inches(3.0)
    if bullets:
        _text(s, Inches(0.8), bottom + Inches(0.14), Inches(11.7), Inches(1.0),
              [(f"•  {b}", 14.5, c, bold) for b, c, bold in bullets], spacing=1.2)
    elif caption:
        _text(s, Inches(0.8), bottom + Inches(0.15), Inches(11.7), Inches(0.8),
              [(caption, 16, ACCENT, True)], align=PP_ALIGN.CENTER)
    return s


def slide_bullets(prs, title, lines, note=None):
    """lines: list of (text, size, colour, bold). Use (" ", 10, INK, False) as a spacer."""
    s = _blank(prs)
    _heading(s, title)
    _text(s, Inches(0.85), Inches(1.6), Inches(11.6), Inches(4.8), list(lines), spacing=1.3)
    if note:
        _text(s, Inches(0.85), Inches(6.35), Inches(11.6), Inches(0.7), [(note, 14, MUTED, False)])
    return s


def slide_two_col(prs, title, left_head, left_items, right_head, right_items, note=None,
                  left_colour=WARN, right_colour=GOOD):
    """items: list of (text, emphasised)."""
    s = _blank(prs)
    _heading(s, title)
    for x, head, items, colour in ((Inches(0.8), left_head, left_items, left_colour),
                                   (Inches(7.0), right_head, right_items, right_colour)):
        _text(s, x, Inches(1.6), Inches(5.5), Inches(0.6), [(head, 22, colour, True)])
        _text(s, x, Inches(2.3), Inches(5.5), Inches(4.0),
              [(f"•  {t}", 15, colour if em else INK, em) for t, em in items], spacing=1.3)
    if note:
        _text(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.7), [(note, 14, MUTED, False)])
    return s


def save_deck(prs, out):
    """Number every slide except the title, then save. Numbering last keeps it right after re-ordering."""
    for i, slide in enumerate(prs.slides, start=1):
        if i > 1:
            _number(slide, i)
    prs.save(out)
    print(f"  wrote {Path(out).name}  ({len(prs.slides)} slides)")
